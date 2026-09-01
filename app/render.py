"""Render a GeneratedResume into a single-page ATS-friendly .docx.

Takes the tailored bullets from a GeneratedResume, joins them back to the
Experience/Project rows they came from (for dates, location, tech stack), and
lays them out with Education and Certifications.

ATS-friendly means: one column, no tables, no text boxes, no images, no content
in headers/footers, standard fonts, real bullet lists, plain section headings.
Everything a parser needs is in the document body in reading order.

Usage:
    python -m app.render --job-id 1 --output /tmp/resume_tiktok.docx
    python -m app.render --job-id 1 --output out.docx --name "Jane Doe" \\
        --email jane@example.com --phone "+61 400 000 000" --link github.com/jane

Contact details come from the Profile singleton row (loaded by app.import_profile
from profile.json's "contact" object). Any flag below overrides the stored value.
"""

import argparse
import json
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from sqlalchemy import select

from app.database import SessionLocal, init_db
from app.models import (
    Certification,
    Education,
    Experience,
    GeneratedResume,
    JobApplication,
    Project,
)
from app.models import Profile as ContactRow

BODY_FONT = "Calibri"
BODY_SIZE = Pt(10)
NAME_SIZE = Pt(18)
HEADING_SIZE = Pt(11)
MARGIN_INCHES = 0.5

# US Letter minus margins, in points — the space a single page actually has.
PAGE_HEIGHT_PT = 11 * 72
USABLE_HEIGHT_PT = PAGE_HEIGHT_PT - 2 * MARGIN_INCHES * 72
USABLE_WIDTH_PT = (8.5 - 2 * MARGIN_INCHES) * 72
# Calibri averages ~0.48em per character across mixed-case prose.
AVG_CHAR_WIDTH_EM = 0.48
LINE_HEIGHT_FACTOR = 1.15


def fmt_date(value: date | None, *, ongoing: str = "Present") -> str:
    return value.strftime("%b %Y") if value else ongoing


def date_range(start: date | None, end: date | None) -> str:
    if not start and not end:
        return ""
    return f"{fmt_date(start, ongoing='')} – {fmt_date(end)}".strip(" –")


def add_horizontal_rule(paragraph) -> None:
    """Bottom border on a paragraph — a rule that parsers ignore safely."""
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "999999")
    borders.append(bottom)
    p_pr.append(borders)


def tighten(paragraph, *, before: int = 0, after: int = 0) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.0


def wrapped_lines(text: str, font_size_pt: float, indent_pt: float = 0) -> int:
    """How many rendered lines `text` takes at this size. Word does the real
    layout; this is close enough to warn about a second page."""
    if not text:
        return 1
    chars_per_line = max(
        20, int((USABLE_WIDTH_PT - indent_pt) / (font_size_pt * AVG_CHAR_WIDTH_EM))
    )
    return max(1, -(-len(text) // chars_per_line))


class ResumeDoc:
    """Thin wrapper so each section reads as layout, not python-docx plumbing.

    Tracks estimated rendered height so we can warn when content will not fit
    on one page — python-docx does no layout, so nothing else knows.
    """

    def __init__(self) -> None:
        self.document = Document()
        self.height_pt = 0.0
        self._setup()

    def _grow(self, text: str, size_pt: float, *, before=0, after=0, indent=0.0) -> None:
        lines = wrapped_lines(text, size_pt, indent)
        self.height_pt += lines * size_pt * LINE_HEIGHT_FACTOR + before + after

    def _setup(self) -> None:
        for section in self.document.sections:
            section.top_margin = section.bottom_margin = Pt(MARGIN_INCHES * 72)
            section.left_margin = section.right_margin = Pt(MARGIN_INCHES * 72)

        normal = self.document.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.font.size = BODY_SIZE
        # East-Asian font mapping, or Word may substitute for some glyphs.
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    def header(self, name: str, contact_parts: list[str]) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tighten(paragraph, after=2)
        run = paragraph.add_run(name)
        run.bold = True
        run.font.size = NAME_SIZE
        self._grow(name, NAME_SIZE.pt, after=2)

        if contact_parts:
            line = self.document.add_paragraph()
            line.alignment = WD_ALIGN_PARAGRAPH.CENTER
            tighten(line, after=6)
            run = line.add_run("  |  ".join(contact_parts))
            run.font.size = Pt(9)
            self._grow("  |  ".join(contact_parts), 9, after=6)

    def section(self, title: str) -> None:
        paragraph = self.document.add_paragraph()
        tighten(paragraph, before=6, after=3)
        run = paragraph.add_run(title.upper())
        run.bold = True
        run.font.size = HEADING_SIZE
        run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        add_horizontal_rule(paragraph)
        self._grow(title, HEADING_SIZE.pt, before=6, after=3)

    def entry(self, left: str, right: str = "", subtitle: str = "") -> None:
        """A bold left-aligned title with an optional right-aligned date.

        If the title and date cannot share a line, the date moves down to the
        subtitle rather than wrapping onto a ragged line of its own.
        """
        char_width = BODY_SIZE.pt * AVG_CHAR_WIDTH_EM
        # Bold runs are wider than the average; leave a gap so they never collide.
        fits = (len(left) * char_width * 1.08 + len(right) * char_width + 24) <= USABLE_WIDTH_PT
        inline_right = right if (right and fits) else ""
        if right and not fits:
            subtitle = f"{subtitle} · {right}" if subtitle else right

        paragraph = self.document.add_paragraph()
        tighten(paragraph, before=3, after=0)
        # Right-align the date with a tab stop rather than a table.
        if inline_right:
            paragraph.paragraph_format.tab_stops.add_tab_stop(
                Pt(USABLE_WIDTH_PT), WD_ALIGN_PARAGRAPH.RIGHT
            )
        run = paragraph.add_run(left)
        run.bold = True
        if inline_right:
            paragraph.add_run("\t" + inline_right)
        self._grow(left, BODY_SIZE.pt, before=3)

        if subtitle:
            sub = self.document.add_paragraph()
            tighten(sub, after=1)
            run = sub.add_run(subtitle)
            run.italic = True
            run.font.size = Pt(9)
            self._grow(subtitle, 9, after=1)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            paragraph = self.document.add_paragraph(item, style="List Bullet")
            tighten(paragraph, after=1)
            paragraph.paragraph_format.left_indent = Pt(12)
            # Bullet glyph plus hanging indent eats ~30pt of the text column.
            self._grow(item, BODY_SIZE.pt, after=1, indent=30)

    def line(self, text: str, *, size: int = 10) -> None:
        paragraph = self.document.add_paragraph()
        tighten(paragraph, after=1)
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        self._grow(text, size, after=1)

    def save(self, path: Path) -> None:
        self.document.save(path)


def load(job_id: int, resume_id: int | None = None) -> dict:
    """Resume content joined back to the profile rows it references."""
    with SessionLocal() as db:
        application = db.get(JobApplication, job_id)
        if application is None:
            raise LookupError(f"no JobApplication with id={job_id}")

        if resume_id is not None:
            resume = db.get(GeneratedResume, resume_id)
            if resume is None or resume.job_application_id != job_id:
                raise LookupError(f"no GeneratedResume id={resume_id} for job {job_id}")
        else:
            resume = db.scalar(
                select(GeneratedResume)
                .where(GeneratedResume.job_application_id == job_id)
                .order_by(GeneratedResume.created_at.desc(), GeneratedResume.id.desc())
                .limit(1)
            )
            if resume is None:
                raise LookupError(
                    f"no GeneratedResume for job {job_id} — run app.generate first"
                )

        content = json.loads(resume.content_json)

        experience_ids = [e["id"] for e in content.get("selected_experiences", [])]
        project_ids = [p["id"] for p in content.get("selected_projects", [])]

        experiences = {
            e.id: e
            for e in db.scalars(select(Experience).where(Experience.id.in_(experience_ids))).all()
        }
        projects = {
            p.id: p for p in db.scalars(select(Project).where(Project.id.in_(project_ids))).all()
        }

        return {
            "application": application,
            "resume_id": resume.id,
            "content": content,
            # Keep the model's ordering; it ranked by relevance.
            "experiences": [
                (entry, experiences.get(entry["id"]))
                for entry in content.get("selected_experiences", [])
            ],
            "projects": [
                (entry, projects.get(entry["id"]))
                for entry in content.get("selected_projects", [])
            ],
            "education": db.scalars(
                select(Education).order_by(Education.end_date.desc().nulls_first())
            ).all(),
            "certifications": db.scalars(
                select(Certification).order_by(Certification.date_earned.desc())
            ).all(),
        }


def build(data: dict, contact: dict) -> ResumeDoc:
    doc = ResumeDoc()

    parts = [
        value
        for value in (contact.get("location"), contact.get("email"), contact.get("phone"))
        if value
    ]
    parts.extend(contact.get("links") or [])
    doc.header(contact.get("name") or "", parts)

    if data["experiences"]:
        doc.section("Experience")
        for entry, row in data["experiences"]:
            company = entry.get("company") or (row.company if row else "")
            title = entry.get("title") or (row.title if row else "")
            dates = date_range(row.start_date, row.end_date) if row else ""
            subtitle_parts = [p for p in ((row.location if row else None),) if p]
            doc.entry(f"{title} — {company}", dates, ", ".join(subtitle_parts))
            doc.bullets(entry.get("bullets") or [])

    if data["projects"]:
        doc.section("Projects")
        for entry, row in data["projects"]:
            name = entry.get("name") or (row.name if row else "")
            tech = ", ".join((row.tech_stack or [])[:6]) if row else ""
            url = (row.github_url or "").replace("https://", "") if row else ""
            doc.entry(name, url, tech)
            doc.bullets(entry.get("bullets") or [])

    if data["education"]:
        doc.section("Education")
        for row in data["education"]:
            degree = row.degree + (f", {row.field}" if row.field else "")
            doc.entry(f"{degree} — {row.institution}", date_range(row.start_date, row.end_date))
            if row.gpa:
                doc.line(f"GPA: {row.gpa}", size=9)

    if data["certifications"]:
        doc.section("Certifications")
        for row in data["certifications"]:
            doc.line(f"{row.name} — {row.issuer} ({fmt_date(row.date_earned, ongoing='')})".strip())

    return doc


def resolve_contact(args) -> dict:
    """Contact details from the Profile row, with CLI flags taking precedence."""
    with SessionLocal() as db:
        row = db.get(ContactRow, 1)

    stored_links = []
    if row:
        stored_links = [
            url for url in (row.github_url, row.linkedin_url, row.portfolio_url) if url
        ]

    return {
        "name": args.name or (row.name if row else None),
        "email": args.email or (row.email if row else None),
        "phone": args.phone or (row.phone if row else None),
        "location": args.location or (row.location if row else None),
        # Any --link replaces the stored set rather than appending to it.
        "links": args.link or stored_links,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Render a generated resume to .docx.")
    parser.add_argument("--job-id", type=int, required=True, help="JobApplication id.")
    parser.add_argument("--output", type=Path, required=True, help="Output .docx path.")
    parser.add_argument("--resume-id", type=int, help="Specific GeneratedResume (default: latest).")
    parser.add_argument("--name", help="Override the stored candidate name.")
    parser.add_argument("--email", help="Override the stored email.")
    parser.add_argument("--phone", help="Override the stored phone.")
    parser.add_argument("--location", help="Override the stored location.")
    parser.add_argument(
        "--link", action="append", help="Link, repeatable. Replaces the stored links."
    )
    args = parser.parse_args()

    init_db()

    try:
        data = load(args.job_id, args.resume_id)
    except LookupError as exc:
        print(f"error: {exc}", file=sys.stderr)
        raise SystemExit(1) from exc

    contact = resolve_contact(args)
    if not contact["name"]:
        print(
            "warning: no candidate name — import a \"contact\" object via app.import_profile, "
            "or pass --name (the resume will have no header name)",
            file=sys.stderr,
        )

    missing = [
        f"{kind} id={entry['id']}"
        for kind, pairs in (("experience", data["experiences"]), ("project", data["projects"]))
        for entry, row in pairs
        if row is None
    ]
    if missing:
        print(
            f"warning: referenced rows no longer exist, rendered without dates: {', '.join(missing)}",
            file=sys.stderr,
        )

    doc = build(data, contact)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)

    application = data["application"]
    print(
        f"Wrote {args.output} — {application.company_name} / {application.role_title} "
        f"(GeneratedResume id={data['resume_id']})"
    )
    print(
        f"  {len(data['experiences'])} experience, {len(data['projects'])} project(s), "
        f"{len(data['education'])} education, {len(data['certifications'])} certification(s)"
    )

    fill = doc.height_pt / USABLE_HEIGHT_PT
    print(f"  estimated page fill: {fill:.0%}")
    if fill > 1.0:
        print(
            f"warning: content is about {fill:.0%} of one page and will likely spill onto a "
            "second — trim bullets, or regenerate with fewer entries",
            file=sys.stderr,
        )


if __name__ == "__main__":
    main()
